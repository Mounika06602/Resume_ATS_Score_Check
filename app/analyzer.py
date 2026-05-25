"""
Resume parsing and AI evaluation engine using LangChain and Google Gemini.

This module provides utility functions for extracting text from multiple document formats
(PDF, DOCX, TXT) and handles the semantic analysis chain that benchmarks
resume content against a target job description.
"""

import os
import pypdf
import docx
from typing import List, Dict, Any
from pydantic import BaseModel, Field
from dotenv import load_dotenv
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain_core.prompts import ChatPromptTemplate
from langchain_core.output_parsers import JsonOutputParser

class ResumeAnalysis(BaseModel):
    """Pydantic schema representing the structured evaluation output of the ATS analysis."""
    match_percentage: int = Field(description="Percentage match between resume and job description (0 to 100)")
    matching_skills: List[str] = Field(description="Key skills and keywords present in both the resume and the job description")
    missing_skills: List[str] = Field(description="Important skills, keywords, or requirements mentioned in the job description but missing or weak in the resume")
    ats_score_factors: Dict[str, int] = Field(description="Simulated ATS scoring breakdown out of 100 (e.g. key: 'Keyword Match', value: score)")
    improvement_suggestions: List[str] = Field(description="Specific, actionable suggestions to improve the resume format, experience description, or keyword inclusion to match the job description")
    recommendation_summary: str = Field(description="A concise summary evaluation of the candidate's fit and general recommendation")

def extract_text_from_pdf(file_stream) -> str:
    """Extract plain text from an in-memory PDF file stream."""
    try:
        reader = pypdf.PdfReader(file_stream)
        text = ""
        # Iterate over all pages and merge text blocks
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
        return text
    except Exception as e:
        raise ValueError(f"Error extracting text from PDF: {str(e)}")

def extract_text_from_docx(file_stream) -> str:
    """Extract plain text from an in-memory DOCX file stream including tables."""
    try:
        doc = docx.Document(file_stream)
        text = ""
        
        # 1. Extract text content from paragraphs
        for para in doc.paragraphs:
            text += para.text + "\n"
            
        # 2. Extract text content from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text]
                if row_text:
                    text += " | ".join(row_text) + "\n"
                    
        return text
    except Exception as e:
        raise ValueError(f"Error extracting text from DOCX: {str(e)}")

def extract_text(file_stream, filename: str) -> str:
    """Detect extension and route file stream to correct parser."""
    ext = os.path.splitext(filename)[1].lower()
    
    if ext == '.pdf':
        return extract_text_from_pdf(file_stream)
    elif ext in ['.docx', '.doc']:
        return extract_text_from_docx(file_stream)
    elif ext in ['.txt', '.md']:
        try:
            # Read plain text files directly by decoding the byte content
            return file_stream.read().decode('utf-8', errors='ignore')
        except Exception as e:
            raise ValueError(f"Error reading text file: {str(e)}")
    else:
        raise ValueError("Unsupported file format. Please upload PDF (.pdf), Word Document (.docx), or Text (.txt).")

def analyze_resume_vs_jd(resume_text: str, jd_text: str, api_key: str = None) -> Dict[str, Any]:
    """
    Compare the resume and job description using Gemini and LangChain.
    Returns a dictionary matching the ResumeAnalysis schema.
    """
    # Force reload variables to capture updates from the web runtime
    load_dotenv(override=True)
    
    # Identify the API Key to use (custom key input override vs environment variable)
    actual_api_key = api_key or os.environ.get("GEMINI_API_KEY")
    
    if not actual_api_key:
        raise ValueError("Google Gemini API Key is missing. Please set GEMINI_API_KEY environment variable or provide it in the web interface.")
        
    # Cascade list of fallback models in case the user has deprecated legacy model definitions in their sandbox
    models_to_try = [
        "gemini-2.5-flash",
        "gemini-2.0-flash",
        "gemini-1.5-flash",
        "gemini-1.5-pro"
    ]
    
    # Instantiate the JSON output parser mapped directly to the Pydantic schema
    parser = JsonOutputParser(pydantic_object=ResumeAnalysis)
    
    # Setup the prompt instruct template
    prompt = ChatPromptTemplate.from_template(
        "You are an expert ATS (Applicant Tracking System) reviewer and hiring manager.\n"
        "Analyze the candidate's Resume against the provided Job Description (JD).\n"
        "Evaluate their alignment, extract matching and missing keywords, perform simulated ATS scoring, "
        "and suggest how the resume can be improved for better matching.\n\n"
        "Resume text:\n{resume_text}\n\n"
        "Job Description:\n{jd_text}\n\n"
        "Ensure you follow the formatting instructions exactly. "
        "Under 'ats_score_factors', provide a dictionary where the keys are short descriptions of the grading metrics "
        "(e.g. 'Keyword Relevance', 'Experience Alignment', 'Education & Certifications', 'Formatting & Structure') "
        "and the values are integer scores (0 to 100) representing how well the resume fares under each category. "
        "The match_percentage should be a fair average or representation of these factors.\n\n"
        "{format_instructions}"
    )
    
    last_error = None
    # Run the model cascading fallback loop
    for model_name in models_to_try:
        try:
            print(f"Attempting analysis using model: {model_name}")
            llm = ChatGoogleGenerativeAI(
                model=model_name,
                google_api_key=actual_api_key,
                temperature=0.1
            )
            
            # Chain configuration: Prompt -> LLM -> JSON Parser
            chain = prompt | llm | parser
            
            result = chain.invoke({
                "resume_text": resume_text,
                "jd_text": jd_text,
                "format_instructions": parser.get_format_instructions()
            })
            print(f"Success: Analysis completed using model: {model_name}")
            return result
        except Exception as e:
            error_str = str(e)
            print(f"Model {model_name} failed with error: {error_str}")
            last_error = e
            
            # If the specific model variant isn't supported or not found in their account tier, move to the next model
            if "not found" in error_str.lower() or "404" in error_str.lower() or "not supported" in error_str.lower():
                continue
            else:
                # Fail immediately for authorization/billing/quota errors to notify the user
                raise RuntimeError(f"Error during AI analysis ({model_name}): {error_str}")
                
    raise RuntimeError(f"All attempted Gemini models failed. Last error: {str(last_error)}")

