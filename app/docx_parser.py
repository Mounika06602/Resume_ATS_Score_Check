"""
Word Document parser service.

This module provides functions for extracting plain text from DOCX files using python-docx.
"""

import docx

def extract_text_from_docx(file_stream) -> str:
    """Extract plain text from an in-memory DOCX file stream including tables."""
    try:
        doc = docx.Document(file_stream)
        text = ""
        
        # 1. Extract text content from paragraphs
        for para in doc.paragraphs:
            text += para.text + "\n"
            
        # 2. Extract text content from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text]
                if row_text:
                    text += " | ".join(row_text) + "\n"
                    
        return text
    except Exception as e:
        raise ValueError(f"Error extracting text from DOCX: {str(e)}")
