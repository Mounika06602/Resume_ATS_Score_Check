import os
import docx

def create_docx_resume():
    doc = docx.Document()
    
    # Title/Header
    title = doc.add_paragraph()
    run = title.add_run("Alex Developer")
    run.bold = True
    run.font.size = docx.shared.Pt(24)
    title.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    
    contact = doc.add_paragraph("alex.developer@email.com | +1-555-0199 | Seattle, WA | github.com/alexdev")
    contact.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    
    # Professional Summary
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph(
        "Detail-oriented Software Engineer with 3+ years of experience specializing in backend web development. "
        "Proficient in Python and Flask, with hands-on experience designing RESTful APIs and optimizing relational database queries."
    )
    
    # Skills
    doc.add_heading("Core Technical Skills", level=1)
    doc.add_paragraph("Languages: Python, JavaScript, SQL, HTML, CSS")
    doc.add_paragraph("Web Frameworks: Flask, Express.js")
    doc.add_paragraph("Databases & Tools: PostgreSQL, SQLite, Git, Linux, Postman")
    
    # Work Experience
    doc.add_heading("Professional Experience", level=1)
    
    # Job 1
    p1 = doc.add_paragraph()
    r1 = p1.add_run("Backend Engineer | Tech Solutions Corp")
    r1.bold = True
    p1.add_run("\nSeattle, WA | Jan 2024 - Present")
    p1.paragraph_format.space_after = docx.shared.Pt(2)
    
    doc.add_paragraph(
        "• Developed and maintained 15+ RESTful API endpoints using Python and Flask, improving mobile app loading times by 25%.\n"
        "• Designed database schemas and optimized complex SQL queries in PostgreSQL, reducing query latency by 40%.\n"
        "• Automated backend unit testing coverage, increasing test reliability from 60% to 92%.\n"
        "• Collaborated with frontend engineers to integrate API responses with responsive React applications."
    )
    
    # Job 2
    p2 = doc.add_paragraph()
    r2 = p2.add_run("Junior Developer | Innovation Labs")
    r2.bold = True
    p2.add_run("\nBoston, MA | Jun 2022 - Dec 2023")
    p2.paragraph_format.space_after = docx.shared.Pt(2)
    
    doc.add_paragraph(
        "• Assisted in building server-side applications using Node.js, Express, and SQLite.\n"
        "• Managed code repositories, branching, and pull requests utilizing Git and GitHub.\n"
        "• Authored technical documentation outlining API structures and setup instructions for new developers."
    )
    
    # Education
    doc.add_heading("Education", level=1)
    p_edu = doc.add_paragraph()
    r_edu = p_edu.add_run("B.S. in Computer Science")
    r_edu.bold = True
    p_edu.add_run("\nUniversity of Washington, Seattle | Graduated 2022")
    
    # Save the file
    doc.save("sample_resume.docx")
    print("Created sample_resume.docx successfully.")

def create_jd_txt():
    jd_content = """Senior Backend Developer - Python & Cloud

We are seeking a Senior Backend Engineer to join our growing product team. In this role, you will lead the architecture, design, and implementation of high-performance microservices.

Key Responsibilities:
- Design, build, and deploy robust APIs and backend microservices using Python and Django.
- Manage data pipelines and cache structures using Redis and PostgreSQL.
- Package services in Docker containers and deploy to AWS using Kubernetes.
- Drive continuous integration and deployment (CI/CD) pipelines to maintain automated releases.
- Mentor junior engineers and conduct technical design and code reviews.

Required Skills & Experience:
- 5+ years of software engineering experience.
- Expert-level knowledge of Python and Django framework (Flask is a plus).
- Deep experience with relational databases (PostgreSQL) and caching (Redis).
- Hands-on cloud experience with AWS (EC2, RDS, S3) and containerization (Docker, Kubernetes).
- Familiarity with CI/CD tools (GitHub Actions, Jenkins).
- Solid understanding of Git, Unix, and software testing practices.
"""
    with open("sample_jd.txt", "w", encoding="utf-8") as f:
        f.write(jd_content)
    print("Created sample_jd.txt successfully.")

if __name__ == "__main__":
    create_docx_resume()
    create_jd_txt()
